"""Office document skills: real Word, PowerPoint, Excel and PDF files."""
import os

from skills.base import skill


def _ensure_dir(path: str) -> None:
    parent = os.path.dirname(os.path.expanduser(path))
    if parent:
        os.makedirs(parent, exist_ok=True)


@skill(
    description="Create a real Word (.docx) document. 'paragraphs' is a list of body "
                "paragraph strings; 'title' is optional.",
    parameters={"type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "title": {"type": "string"},
                    "paragraphs": {"type": "array", "items": {"type": "string"}},
                }, "required": ["path"]},
)
def create_word_document(path: str, title: str = "", paragraphs: list = None) -> str:
    from docx import Document
    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    for p in (paragraphs or []):
        doc.add_paragraph(str(p))
    _ensure_dir(path)
    doc.save(os.path.expanduser(path))
    return f"Word document created: {path}"


@skill(
    description="Create a real PowerPoint (.pptx). 'slides' is a list of objects like "
                "{'title': 'Slide Title', 'content': 'one bullet per line'}.",
    parameters={"type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "slides": {"type": "array", "items": {
                        "type": "object",
                        "properties": {"title": {"type": "string"},
                                       "content": {"type": "string"}},
                    }},
                }, "required": ["path", "slides"]},
)
def create_powerpoint(path: str, slides: list) -> str:
    from pptx import Presentation
    prs = Presentation()
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(s.get("title", ""))
        content = str(s.get("content", "")).strip()
        lines = content.splitlines() if content else [""]
        body = slide.placeholders[1].text_frame
        body.text = lines[0]
        for line in lines[1:]:
            body.add_paragraph().text = line
    _ensure_dir(path)
    prs.save(os.path.expanduser(path))
    return f"PowerPoint created: {path}"


def _coerce(cell):
    s = str(cell).strip()
    for cast in (int, float):
        try:
            return cast(s)
        except ValueError:
            pass
    return str(cell)


@skill(
    description="Create a real Excel (.xlsx) file. 'rows' is a list of rows (first row "
                "usually headers); numeric-looking cells become real numbers.",
    parameters={"type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "sheet_name": {"type": "string"},
                    "rows": {"type": "array", "items": {"type": "array"}},
                }, "required": ["path", "rows"]},
)
def create_excel(path: str, rows: list, sheet_name: str = "Sheet1") -> str:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name or "Sheet1"
    for row in rows:
        ws.append([_coerce(c) for c in row])
    _ensure_dir(path)
    wb.save(os.path.expanduser(path))
    return f"Excel file created: {path}"


@skill(
    description="Extract text from a PDF file (first pages). Use to READ PDFs for the user.",
    parameters={"type": "object",
                "properties": {"path": {"type": "string"},
                               "max_pages": {"type": "integer"}},
                "required": ["path"]},
)
def read_pdf(path: str, max_pages: int = 5) -> str:
    from pypdf import PdfReader
    reader = PdfReader(os.path.expanduser(path))
    pages = reader.pages[:max(1, min(int(max_pages), 20))]
    text = "\n".join((p.extract_text() or "") for p in pages)
    return f"PDF has {len(reader.pages)} pages. Text of first {len(pages)}:\n{text[:4000]}"


@skill(
    description="Create a simple PDF with ENGLISH (Latin) text. 'lines' is a list of "
                "paragraph strings. (Persian PDFs are not supported yet.)",
    parameters={"type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "title": {"type": "string"},
                    "lines": {"type": "array", "items": {"type": "string"}},
                }, "required": ["path"]},
)
def write_pdf(path: str, title: str = "", lines: list = None) -> str:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    if title:
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, title)
        pdf.ln(12)
    pdf.set_font("Helvetica", size=12)
    for line in (lines or []):
        pdf.multi_cell(0, 8, str(line))
        pdf.ln(2)
    _ensure_dir(path)
    pdf.output(os.path.expanduser(path))
    return f"PDF created: {path}"